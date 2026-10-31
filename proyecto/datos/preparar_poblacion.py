import csv,sys
from openpyxl import load_workbook
src,out=sys.argv[1:]
w=load_workbook(src,read_only=True,data_only=True)
s=w['PobMunicipalxÁreaSexoEdad']
with open(out,'w',newline='',encoding='utf-8') as f:
 cols=['codigo_municipio','municipio','anio','poblacion_total','hombres','mujeres','pob_18_44','pob_18_64','hombres_18_44','mujeres_18_44']
 writer=csv.DictWriter(f,fieldnames=cols,lineterminator="\n");writer.writeheader();n=0
 for r in s.iter_rows(min_row=10,values_only=True):
  if str(r[0]).zfill(2)!='15' or r[5]!='Total' or r[4] not in (2020,2021,2022,2023,2024,2025): continue
  h=sum(r[9+18:9+45]); m=sum(r[110+18:110+45]); h64=sum(r[9+18:9+65]); m64=sum(r[110+18:110+65])
  writer.writerow(dict(codigo_municipio=str(r[2]).zfill(5),municipio=r[3],anio=r[4],poblacion_total=r[6],hombres=r[7],mujeres=r[8],pob_18_44=h+m,pob_18_64=h64+m64,hombres_18_44=h,mujeres_18_44=m));n+=1
print(n)
